from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Unity\TSS2\Layer_2_Switch_CLI_Programming_Review.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5), ("Heading 3", 10.5)):
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles[name].font.bold = True
styles["Title"].paragraph_format.space_after = Pt(4)
title_style_pPr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_pPr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_pPr.remove(title_style_border)
styles["Heading 1"].paragraph_format.space_before = Pt(13)
styles["Heading 1"].paragraph_format.space_after = Pt(5)
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(3)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table, color="D9D9D9"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    border = tblPr.first_child_found_in("w:tblBorders")
    if border is None:
        border = OxmlElement("w:tblBorders")
        tblPr.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = border.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            border.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)

def set_cell(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    borders(t)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, (255, 255, 255), 8.5)
        shade(t.rows[0].cells[i], "1F4E78")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value, False, None, 8.3)
            if ri % 2 == 1:
                shade(cells[i], "F2F6FA")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p

title_p = doc.add_paragraph("Layer 2 Switch CLI Programming Review", style="Title")
title_pPr = title_p._p.get_or_add_pPr()
for border_tag in ("w:pBdr", "w:bottom", "w:top", "w:left", "w:right"):
    border = title_pPr.find(qn(border_tag))
    if border is not None:
        title_pPr.remove(border)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(9)
r = p.add_run("Supported commands, runtime state effects, and incompleteness findings")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(68, 68, 68)

doc.add_paragraph("Purpose and conclusion", style="Heading 1")
doc.add_paragraph(
    "This review covers the current Layer 2 switch CLI implementation in TSS2. It maps the supported command groups to the records they modify and identifies behavior that is incomplete, misleading, or not yet connected to the simulated network state. The main MVP workflow is present: access VLANs, trunks, allowed and native VLANs, interface shutdown, VLAN creation, show commands, and configuration save."
)
doc.add_paragraph(
    "The most important gaps are default trunk behavior, suspended VLAN presentation, the absence of a real MAC address table, and the fact that switch devices currently accept router-style IP and static-route commands."
)
doc.add_paragraph("Scope and source of truth", style="Heading 2")
bullet("Primary command implementation: Assets/Scripts/TssCliCommandProcessor.cs")
bullet("Runtime state records: Assets/Scripts/TssNetworkSessionState.cs")
bullet("Change propagation: Assets/Scripts/TssTrainingSession.cs")
bullet("Visible terminal path: Assets/Scripts/TssTerminalSessionPresenter.cs")
bullet("Existing command coverage: Assets/Tests/EditMode/TssDeviceSessionTests.cs")

doc.add_paragraph("Command support by mode", style="Heading 1")

doc.add_paragraph("EXEC and privileged EXEC", style="Heading 2")
table(["Command group", "Support", "State effect"], [
    ("enable, en", "Supported", "Changes NetworkDeviceState.Mode; may set AwaitingEnablePassword."),
    ("disable", "Supported", "Changes mode to UserExec."),
    ("configure terminal, conf t", "Supported", "Changes mode to GlobalConfig."),
    ("exit, end", "Supported", "Changes CLI mode only."),
    ("write memory, copy run start", "Supported", "Clones state into StartupSnapshot and records save evidence."),
    ("show running-config", "Supported", "Reads current device state."),
    ("show startup-config", "Supported", "Reads StartupSnapshot."),
    ("show vlan, show vlan brief", "Supported", "Reads VLAN records and records VLAN evidence."),
    ("show interfaces trunk", "Supported", "Reads trunk records and records trunk evidence."),
    ("show interfaces, show ip interface brief", "Supported", "Reads interface records."),
    ("show mac address-table", "Partial", "Returns a fixed no-dynamic-entries response; no MAC table state exists."),
    ("show ip route", "Supported but questionable", "Reads NetworkRouteState; static routes can be configured on switches."),
    ("ping", "Placeholder", "Always reports 0 percent success; does not evaluate network state."),
    ("clear", "UI-only", "Clears terminal output, not network state."),
], [1.65, 1.0, 4.35])

doc.add_paragraph("Global configuration", style="Heading 2")
table(["Command", "State record changed"], [
    ("hostname, no hostname", "Hostname; also attempts to rename the installed scene or endpoint record."),
    ("enable password, no enable password", "EnablePassword."),
    ("enable secret, no enable secret", "EnableSecret."),
    ("service password-encryption", "PasswordEncryption; affects rendering, not stored-value encryption."),
    ("ip domain-lookup, no ip domain-lookup", "DomainLookupDisabled."),
    ("banner motd, no banner motd", "Motd."),
    ("line console 0", "Selects console and enters LineConfig."),
    ("line vty 0 4", "Selects VTY and enters LineConfig."),
    ("vlan <id>", "Creates or selects Vlans[id] and enters VlanConfig."),
    ("no vlan <id>", "Removes a VLAN database entry, except VLAN 1."),
    ("ip route, no ip route", "Adds or removes NetworkRouteState records."),
    ("interface, interface range", "Selects or creates NetworkInterfaceState records."),
], [2.15, 4.85])

doc.add_paragraph("Access interface commands", style="Heading 2")
table(["Command", "State effect"], [
    ("switchport mode access", "IsTrunk=false; SwitchportModeConfigured=true."),
    ("no switchport mode access", "IsTrunk=false; SwitchportModeConfigured=false."),
    ("switchport access vlan <id>", "Sets AccessVlan and SwitchportAccessVlanConfigured."),
    ("no switchport access vlan", "Resets AccessVlan to VLAN 1 and clears the configured flag."),
    ("shutdown / no shutdown", "Sets AdminUp false or true; state change refreshes carriers and endpoint connectivity."),
    ("ip address <address> <mask>", "Sets Address and Mask, including on switch physical interfaces."),
    ("no ip address", "Clears Address and Mask."),
], [2.15, 4.85])

doc.add_paragraph("Trunk interface commands", style="Heading 2")
table(["Command", "State effect"], [
    ("switchport trunk encapsulation dot1q", "Sets TrunkEncapsulationDot1q=true."),
    ("no switchport trunk encapsulation dot1q", "Sets TrunkEncapsulationDot1q=false."),
    ("switchport mode trunk", "IsTrunk=true; SwitchportModeConfigured=true."),
    ("no switchport mode trunk", "IsTrunk=false; SwitchportModeConfigured=false."),
    ("switchport trunk allowed vlan <list>", "Replaces AllowedVlans with the supplied VLAN IDs."),
    ("no switchport trunk allowed vlan", "Clears AllowedVlans."),
    ("switchport trunk native vlan <id>", "Sets NativeVlan."),
    ("no switchport trunk native vlan", "Resets NativeVlan to VLAN 1."),
], [2.65, 4.35])

doc.add_paragraph("Line and VLAN configuration", style="Heading 2")
table(["Mode", "Commands", "State effect"], [
    ("LineConfig", "password, no password", "Updates ConsolePassword or VtyPassword."),
    ("LineConfig", "login, no login", "Updates ConsoleLoginRequired or VtyLoginRequired."),
    ("VlanConfig", "name, no name", "Stores the VLAN display name in Vlans[id]."),
    ("VlanConfig", "state suspend, no state suspend", "Appends or removes '(suspended)' inside the VLAN name string."),
], [1.15, 2.25, 4.2])

doc.add_paragraph("Runtime state behavior", style="Heading 1")
doc.add_paragraph(
    "Mutating commands generally pass through Changed(), which calls TssTrainingSession.NotifyNetworkStateChanged(). That refreshes physical carriers, placed desktop connectivity, and the session StateChanged event. Read-only show commands do not call the notification path. Some show commands still mutate evidence sets directly."
)
doc.add_paragraph("Layer 2 availability", style="Heading 2")
doc.add_paragraph(
    "For switch interfaces, the runtime evaluator requires the interface to be administratively up, physically up, and attached to an active VLAN. Access ports use AccessVlan. Trunks require at least one active VLAN in AllowedVlans. VLAN suspension is inferred from the text stored in the VLAN name."
)

doc.add_paragraph("Incompleteness findings", style="Heading 1")
table(["Priority", "Finding", "Why it matters"], [
    ("High", "New trunks have no allowed VLANs by default.", "AllowedVlans starts empty, and Layer 2 availability requires Any(...). A trunk created with only switchport mode trunk cannot carry traffic until an allowed list is explicitly configured."),
    ("High", "show vlan always reports active.", "The presenter strips '(suspended)' and hardcodes active, so state suspend is not visible in the primary inspection command."),
    ("High", "Switches accept router-style IP and static-route commands.", "ip address and ip route are stored on switch state even though no switch management SVI model exists. This blurs Layer 2 and Layer 3 behavior."),
    ("Medium", "show interfaces trunk does not filter operational status.", "It reports every IsTrunk interface, including administratively down or carrier-down interfaces, despite describing active trunks."),
    ("Medium", "No dynamic MAC address-table state exists.", "show mac address-table cannot display learned MAC addresses, VLANs, or ports."),
    ("Medium", "ping is not connected to network state.", "The command always returns failure and cannot validate switch configuration."),
    ("Medium", "VLAN suspension is encoded in the VLAN name.", "Repeated suspend commands can duplicate the suffix, and status is not represented as a typed field."),
    ("Low", "Help output is incomplete.", "The detailed command reference includes more commands than the mode-specific '?' output, especially no forms and password or banner commands."),
    ("Low", "CLI mode changes do not notify the training session.", "The terminal updates from the returned result, but other state listeners do not receive mode-only changes."),
], [0.7, 2.25, 4.95])

doc.add_paragraph("Recommended review order", style="Heading 1")
for text in [
    "Decide whether an empty trunk allowed list means all VLANs or no VLANs in this training simulator.",
    "Make show vlan display suspended status from the same state used by the connectivity evaluator.",
    "Decide whether Layer 2 switches should reject ip route and physical-interface ip address, or add an explicit management SVI model.",
    "Filter show interfaces trunk by AdminUp and CarrierUp when the output is intended to show active trunks.",
    "Leave MAC learning and functional ping deferred unless they are required for the current training scenario.",
]:
    bullet(text)

doc.add_paragraph("Bottom line", style="Heading 1")
doc.add_paragraph(
    "The current implementation supports the core switch programming loop: access VLAN assignment, trunk configuration, VLAN creation and naming, interface shutdown, operational inspection, and configuration save. It is sufficient for the existing MVP workflow, but the five high-value review items above should be resolved or explicitly documented before treating the switch CLI as complete."
)

doc.add_paragraph("Review basis", style="Heading 2")
doc.add_paragraph("Reviewed October 20 2026 against the current TSS2 source and EditMode command coverage. This document is a source and state-model review; it does not claim a new Play Mode smoke test was run for every command listed.")

doc.save(OUT)
print(OUT)
