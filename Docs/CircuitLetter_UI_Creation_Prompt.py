from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Unity\TSS2")
OUTPUT = ROOT / "Docs" / "CircuitLetter_UI_Creation_Prompt.docx"
IMAGE = ROOT / "Docs" / "Images" / "binder_image" / "Circuit Letter.png"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def borders(cell, color="D9D9D9", size="6"):
    props = cell._tc.get_or_add_tcPr()
    border = props.first_child_found_in("w:tcBorders")
    if border is None:
        border = OxmlElement("w:tcBorders")
        props.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = border.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            border.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
    styles[style_name].font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    styles[style_name].font.size = Pt(size)
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Circuit Letter UI Creation Prompt")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
run = subtitle.add_run("TSS2 Binder Reference implementation brief")
run.italic = True
run.font.color.rgb = RGBColor(75, 85, 99)

intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "Use this prompt to begin creating and reviewing the Circuit Letter tab in Unity. "
    "The goal is to preserve the visual language of the supplied reference while making the tab "
    "scenario-driven, editable, and safe for future training scenarios."
)

add_heading(doc, "Reference image", 1)
image_para = doc.add_paragraph()
image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_para.add_run().add_picture(str(IMAGE), width=Inches(6.8))
caption = doc.add_paragraph("Circuit Letter reference image", style="Caption")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_heading(doc, "Implementation prompt", 1)
prompt = (
    "Create or refine the Circuit Letter Binder Reference UI in the Unity TSS2 project. "
    "Use Assets/Prefabs/UI/TssCircuitLetterTab.prefab as the target prefab and the supplied Circuit Letter image "
    "as the visual reference. Keep the tab connected to the Binder Reference flow in TssRuntimeUi and keep the "
    "scenario as the source of authored content. Do not create a DOCX or duplicate scenario data inside the prefab."
)
doc.add_paragraph(prompt)

add_heading(doc, "Required visual structure", 1)
for item in [
    "Keep the existing Binder tab frame and dark technical visual style.",
    "Expose and preserve the hierarchy objects SectionTitle, ScenarioHeading, IntroParagraph, Table, and Note.",
    "Use a clear Circuit Handoff Letter heading in the same blue technical accent used by the reference image.",
    "Keep ScenarioHeading available for the scenario name or identifier, for example CIRCUIT LETTER - {ScenarioId}.",
    "Keep IntroParagraph above the circuit block and allow it to wrap without clipping.",
    "Use the bordered monospace circuit block for the generated circuit rows. Preserve readable spacing between label and value columns.",
    "Keep Note below the circuit block for the customer default-route instruction.",
    "Match the existing Binder Reference typography, row heights, padding, and responsive layout conventions; do not solve overflow by making the text unreadably small.",
]:
    add_bullet(doc, item)

add_heading(doc, "Scenario and runtime behavior", 1)
for item in [
    "Expose SectionTitle, ScenarioHeading, IntroParagraph, and Note through the scenario Binder Reference definition.",
    "Keep the circuit table rows runtime-driven from TssNetworkRunProfile. Do not hard-code generated circuit values in the scenario asset.",
    "Use scenario content for authored wording and support tokens such as {ScenarioId}, {CircuitId}, {WanProviderAddress}, and {WanCustomerAddress}.",
    "Resolve the customer/provider addresses and circuit ID from the active training run; the UI must update when the scenario or run profile changes.",
    "Keep the existing fallback behavior for scenarios that do not yet define Circuit Letter content.",
    "Preserve the existing SectionRow05 and TssBinderPanel prefab connection; do not replace the runtime Binder tab routing with a second lookup path.",
]:
    add_bullet(doc, item)

add_heading(doc, "Expected scenario fields", 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = (Inches(1.65), Inches(2.15), Inches(3.0))
for i, width in enumerate(widths):
    table.columns[i].width = width
headers = ("Field", "Purpose", "Example")
for cell, value in zip(table.rows[0].cells, headers):
    shade(cell, "1F4E78")
    set_cell_text(cell, value, bold=True, color="FFFFFF")
rows = [
    ("sectionTitle", "Binder section label", "Circuit Letter"),
    ("scenarioHeading", "Scenario-specific heading", "CIRCUIT LETTER - {ScenarioId}"),
    ("intro", "Authored context above the circuit block", "Tapped to the NID on the IT room wall..."),
    ("note", "Authored instruction below the circuit block", "Customer shall install a default route toward {WanProviderAddress}."),
]
for row_index, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        shade(cells[i], "F4F7FA" if row_index % 2 == 0 else "FFFFFF")
        set_cell_text(cells[i], value)

add_heading(doc, "Acceptance checks", 1)
for item in [
    "The prefab contains SectionTitle, ScenarioHeading, IntroParagraph, Table, and Note with no missing references.",
    "The tab appears in the Binder Reference at SectionRow05 and opens without duplicate rows or stale content.",
    "A loaded scenario supplies the authored title, heading, introduction, and note.",
    "The active run profile supplies the circuit ID and WAN addresses dynamically.",
    "The generated values update when a different scenario/run seed is loaded.",
    "Long scenario text wraps cleanly and does not overlap the table or note.",
    "Unity compiles with no errors, and Play Mode shows the updated Circuit Letter tab through the real Binder workflow.",
]:
    add_bullet(doc, item)

add_heading(doc, "Relevant project paths", 1)
for path in [
    "Assets/Prefabs/UI/TssCircuitLetterTab.prefab",
    "Assets/Prefabs/UI/TssBinderPanel.prefab",
    "Assets/Scripts/TssRuntimeUi.cs",
    "Assets/Scripts/TssBinderReferenceTabPresenter.cs",
    "Assets/Scripts/TssScenarioDefinition.cs",
    "Assets/Data/TSS/Scenarios/Balanced_12_Piece_Lab.asset",
]:
    add_bullet(doc, path)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("TSS2 Circuit Letter UI creation brief")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUTPUT)
print(OUTPUT)
