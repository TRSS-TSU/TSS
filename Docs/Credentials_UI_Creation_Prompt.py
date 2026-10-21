from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Unity\TSS2")
OUTPUT = ROOT / "Docs" / "Credentials_UI_Creation_Prompt.docx"
IMAGE = ROOT / "Docs" / "Images" / "binder_image" / "Credentials.png"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def borders(cell, color="D9D9D9", size="6"):
    props = cell._tc.get_or_add_tcPr()
    border = props.first_child_found_in("w:tcBorders")
    if border is None:
        border = OxmlElement("w:tcBorders")
        props.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = border.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            border.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
    styles[style_name].font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    styles[style_name].font.size = Pt(size)
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.add_run("Credentials UI Creation Prompt")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
run = subtitle.add_run("TSS2 Binder Reference implementation brief")
run.italic = True
run.font.color.rgb = RGBColor(75, 85, 99)

intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "Use this prompt to begin creating and reviewing the Credentials tab in Unity. "
    "The objective is a readable credentials and hardening sheet that matches the supplied reference, "
    "keeps sensitive-looking values in the generated training profile, and keeps instructor-authored guidance in the scenario Binder Reference."
)

add_heading(doc, "Reference image", 1)
image_para = doc.add_paragraph()
image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
image_para.add_run().add_picture(str(IMAGE), width=Inches(6.8))
caption = doc.add_paragraph("Credentials reference image", style="Caption")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_heading(doc, "Implementation prompt", 1)
doc.add_paragraph(
    "Create or refine the Credentials Binder Reference UI in the Unity TSS2 project. "
    "Use Assets/Prefabs/UI/TssCredentialsTab.prefab as the target prefab and the supplied Credentials image as the visual reference. "
    "Keep the tab connected to the Binder Reference flow in TssRuntimeUi and keep the scenario as the source of authored content. "
    "Do not duplicate generated credentials inside the prefab or create a second runtime data path."
)

add_heading(doc, "Required visual structure", 1)
for item in [
    "Keep the existing dark technical Binder visual language and tab placement.",
    "Expose and preserve the hierarchy objects SectionTitle, ScenarioHeading, IntroParagraph, Table, and Note.",
    "Use a clear blue technical heading such as CREDENTIALS & HARDENING.",
    "Keep ScenarioHeading available for the scenario name or identifier, for example CREDENTIALS - {ScenarioId}.",
    "Place IntroParagraph above the two-column credentials table and allow it to wrap cleanly.",
    "Use an ITEM / VALUE header row with distinct label and value columns.",
    "Use the established amber or gold accent for item labels and readable light text for values.",
    "Keep Note beneath the table as a concise security or hardening instruction. It should support multiple lines without clipping.",
    "Preserve existing Binder spacing, font sizing, row heights, padding, and responsive behavior.",
]:
    add_bullet(doc, item)

add_heading(doc, "Scenario and runtime behavior", 1)
for item in [
    "Expose SectionTitle, ScenarioHeading, IntroParagraph, and Note through the scenario Binder Reference definition.",
    "Keep the credentials table values runtime-driven from TssNetworkRunProfile. Do not hard-code generated passwords or secrets in the scenario asset.",
    "Use scenario content for authored wording and support {ScenarioId} in the scenario heading plus profile tokens such as {ConsolePassword}, {VtyPassword}, {EnableSecret}, and {Motd} where needed.",
    "Resolve generated credential values from the active training run so the UI changes with the selected scenario/run seed.",
    "Keep non-secret instructional values such as the local user, password-encryption command, and domain-lookup command explicit and readable.",
    "Keep the existing fallback behavior for scenarios that do not yet define Credentials Binder Reference content.",
    "Preserve the existing SectionRow07 and TssBinderPanel prefab connection and the current presenter routing.",
]:
    add_bullet(doc, item)

add_heading(doc, "Expected scenario fields", 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for i, width in enumerate((Inches(1.65), Inches(2.15), Inches(3.0))):
    table.columns[i].width = width
for cell, value in zip(table.rows[0].cells, ("Field", "Purpose", "Example")):
    shade(cell, "1F4E78")
    set_cell_text(cell, value, bold=True, color="FFFFFF")
rows = [
    ("sectionTitle", "Binder section label", "Credentials"),
    ("scenarioHeading", "Scenario-specific heading", "CREDENTIALS - {ScenarioId}"),
    ("intro", "Authored guidance above the table", "Apply identically to AN-1, AN-2, and CN-1."),
    ("note", "Authored security guidance below the table", "The banner is a legal requirement, not decoration."),
]
for row_index, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        shade(cells[i], "F4F7FA" if row_index % 2 == 0 else "FFFFFF")
        set_cell_text(cells[i], value)

add_heading(doc, "Generated and fixed credential rows", 1)
for item in [
    "Console password: generated profile console password",
    "VTY password: generated profile VTY password",
    "Enable secret: generated profile enable secret",
    "Local user: admin / 12#$Qwert",
    "Login banner: generated profile MOTD",
    "Password encryption: service password-encryption",
    "Domain lookup: no ip domain-lookup",
]:
    add_bullet(doc, item)

add_heading(doc, "Acceptance checks", 1)
for item in [
    "The prefab contains SectionTitle, ScenarioHeading, IntroParagraph, Table, and Note with no missing references.",
    "The tab appears in the Binder Reference at SectionRow07 and opens without duplicate rows or stale content.",
    "A loaded scenario supplies the authored title, heading, introduction, and note.",
    "The active run profile supplies the console password, VTY password, enable secret, and MOTD dynamically.",
    "Changing the scenario/run seed changes generated credential values without changing authored guidance.",
    "Sensitive-looking generated values remain readable without being truncated or visually confused with labels.",
    "Long intro or note text wraps cleanly and does not overlap the table.",
    "Unity compiles with no errors, and Play Mode shows the updated Credentials tab through the real Binder workflow.",
]:
    add_bullet(doc, item)

add_heading(doc, "Relevant project paths", 1)
for path in [
    "Assets/Prefabs/UI/TssCredentialsTab.prefab",
    "Assets/Prefabs/UI/TssBinderPanel.prefab",
    "Assets/Scripts/TssRuntimeUi.cs",
    "Assets/Scripts/TssBinderReferenceTabPresenter.cs",
    "Assets/Scripts/TssScenarioDefinition.cs",
    "Assets/Data/TSS/Scenarios/Balanced_12_Piece_Lab.asset",
]:
    add_bullet(doc, path)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("TSS2 Credentials UI creation brief")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUTPUT)
print(OUTPUT)
