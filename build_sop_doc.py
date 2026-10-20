from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Unity\TSS2\SOP_Panel_Tab_Content.docx"

tabs = [
    ("Rack Build Order", """GROUND FLOOR / RACK01
Top to bottom: PDU1, Patch1, AN-1, CN-1, IT2 (server).

SECOND FLOOR / RACK02
Top to bottom: PDU2, Patch2, AN-2.

WORK AREAS
IT1: IT room | Reception1 and Reception2: Reception
Data1: Office 1 | Data2 and Data3: Office 2
AP-1 and TestLaptop: Guest Lounge. Use these exact names."""),
    ("Layer 2 Plan", """Create on both switches:
Management: VLAN 390
Reception: VLAN 180
Data: VLAN 230
Native: VLAN 330 (native)
Dead_Vlan: VLAN 888 (unused ports)

ACCESS PORTS
AN-1 Fa0/2 IT1; Fa0/3 Reception1; Fa0/4 Reception2; Fa0/5 IT2
AN-2 Fa0/2 Data1; Fa0/3 Data2; Fa0/4 Data3

802.1Q TRUNKS
AN-1 Fa0/1 to AN-2 Fa0/1; AN-1 Fa0/6 to CN-1 Fa0/0
Set dot1q, trunk mode, native VLAN, and allow Management, Reception, and Data."""),
    ("Addressing and Gateways", """MANAGEMENT / VLAN 390
Network 192.168.175.0/24 | gateway 192.168.175.1
RECEPTION / VLAN 180
Network 192.168.17.0/24 | gateway 192.168.17.1
DATA / VLAN 230
Network 192.168.54.0/24 | gateway 192.168.54.1

ENDPOINT ADDRESSES
IT1 192.168.175.11  IT2 192.168.175.12
Reception1 192.168.17.21  Reception2 192.168.17.22
Data1 192.168.54.31  Data2 192.168.54.32  Data3 192.168.54.33
AP-1 192.168.17.50 | TestLaptop 192.168.17.60 (DHCP)

CN-1: Fa0/0.<VLAN> subinterfaces use dot1q and the matching gateway above."""),
    ("Cable Schedule", """INTER-CLOSET LINK
Rack01 Patch1 Port48 to Rack02 Patch2 Port48: permanent punch-down run; complete both terminations.
Patch cords: AN-1 Fa0/1 to Patch1 Port48: crossover; AN-2 Fa0/1 to Patch2 Port48: straight-through.
CN-1 Fa0/0 to AN-1 Fa0/6: straight-through.

DEMARC
CN-1 Fa0/1 to Patch1 Port18: straight-through.

GROUND FLOOR
IT1 to AN-1 Fa0/2; IT2 to AN-1 Fa0/5
Reception1/2 to wall ports 1/2, then Patch1 ports 1/2 to AN-1 Fa0/3/4.

SECOND FLOOR
Data1/2/3 to Office wall ports, then Patch2 ports 1/5/6 to AN-2 Fa0/2/3/4.

CIRCUIT AND WIRELESS
CN-1 Fa0/1 to Patch1 Port18; AP-1 to any Reception Wall Port 01-04, then the matching Patch1 Port 01-04 to AN-1 Fa0/7.
Complete each punch-down and patch cord in the structured path."""),
    ("Carrier Circuit Letter", """CIRCUIT ID
CIR-599795

CUSTOMER / CN-1 Fa0/1
Address 10.197.118.1 /30
Provider next hop 10.197.118.2

PHYSICAL HANDOFF
Patch CN-1 Fa0/1 to Rack01 Patch1 Port18. The carrier side is landed at the DEMARC.

ROUTING
Configure a default route to 10.197.118.2. Verify with show ip route and an endpoint ping."""),
    ("Guest Wireless", """AP-1 MANAGEMENT
Address 192.168.17.50 on the Reception network.

GUEST WI-FI
SSID TSS-GUEST-986
Passphrase [REDACTED_SECRET]
Channel 11; disable WPS and apply the secured settings.

UPLINK
AP-1 to Reception Wall Port 01-04, then the matching Rack01 Patch1 Port 01-04 to AN-1 Fa0/7 in the Reception VLAN.

TEST LAPTOP
Join the guest SSID, use DHCP, then verify 192.168.17.60, provider reachability, and the RedWolf page."""),
    ("Run Credentials", """NETWORK DEVICES
Enable secret: [REDACTED_SECRET]
Console password: [REDACTED_SECRET]
VTY password: [REDACTED_SECRET]

ACCESS POINT
Administrator password: [REDACTED_SECRET]

These credentials are generated for this run. Keep them in the binder; do not copy them into notes or debug output."""),
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for style_name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 11)]:
    st = styles[style_name]
    st.font.name = 'Aptos Display' if style_name == 'Title' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), st.font.name)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('SOP Panel Tab Content')

p = doc.add_paragraph()
p.add_run('Purpose. ').bold = True
p.add_run('This document contains the information shown in each tab of the Unity SOP Panel, captured from the live training scenario for review and editing. The seven tabs are kept in panel order. Generated credentials and the Wi-Fi passphrase are intentionally redacted in this editable review copy.')

doc.add_paragraph('Source: live SOP Panel runtime state in the TSS2 Unity project.', style='Subtitle')

for i, (name, body) in enumerate(tabs, 1):
    if i > 1:
        doc.add_page_break()
    doc.add_heading(f'{i}. {name}', level=1)
    for block in body.split('\n\n'):
        p = doc.add_paragraph()
        lines = block.split('\n')
        for j, line in enumerate(lines):
            if j:
                p.add_run().add_break()
            if line.isupper() and not line.startswith('AN-'):
                r = p.add_run(line)
                r.bold = True
            else:
                p.add_run(line)

doc.save(OUT)
print(OUT)
