from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Unity\TSS2\Customer_Switch_and_Router_CLI_Capabilities_Report.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tbl_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        shade(cell, "1F4E78")
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=9)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            if row_index % 2:
                shade(cells[i], "F2F6FA")
            set_cell_text(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders_el = p_pr.find(qn("w:pBdr"))
    if borders_el is not None:
        p_pr.remove(borders_el)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
title_style_ppr = styles["Title"]._element.pPr
if title_style_ppr is not None:
    title_style_borders = title_style_ppr.find(qn("w:pBdr"))
    if title_style_borders is not None:
        title_style_ppr.remove(title_style_borders)
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.add_run("Customer Report Switch and Router CLI Capabilities")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("Current tested functions and operating limits")
run.italic = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(89, 89, 89)

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "The current switch and router command-line interfaces support the core configuration and inspection activities needed for network training. "
    "The tested workflow includes entering configuration modes, setting hostnames and security-related options, configuring interfaces, creating VLANs, building an 802.1Q trunk, inspecting the resulting state, and saving the running configuration."
)
doc.add_paragraph(
    "The switch path is the more complete of the two: access VLANs, the trunk, native VLAN selection, allowed VLANs, inspection commands, and configuration save behavior were all accepted and verified. "
    "The router path supports interface addressing, 802.1Q subinterface configuration, static routing, inspection, and saving configuration. Router subinterface context handling remains an operational limitation: the tested workflow requires explicit interface exits, and retention of subinterface state was not treated as a clean pass in the source test run."
)

doc.add_heading("Capability Overview", level=1)
add_table(doc,
    ["Area", "Switch CLI", "Router CLI"],
    [
        ("Command modes", "User EXEC, Privileged EXEC, Global Configuration, Interface Configuration", "User EXEC, Privileged EXEC, Global Configuration, Interface Configuration"),
        ("Identity and basic settings", "Hostname, DNS lookup setting, password-encryption setting, enable secret", "Hostname, DNS lookup setting, password-encryption setting, enable secret"),
        ("Interface configuration", "Interface selection, shutdown state, access VLAN assignment, trunk settings", "Physical interface selection, no shutdown, interface addressing"),
        ("Layer 2 services", "VLAN creation and naming; access and trunk behavior", "802.1Q subinterface encapsulation path exercised"),
        ("Layer 3 services", "Not the primary tested function", "IP addressing and static route configuration"),
        ("Inspection", "VLAN and trunk inspection", "Running configuration and route inspection"),
        ("Persistence", "Running configuration saved to startup state", "Running configuration saved to startup state"),
    ], [1.45, 2.65, 2.65])

doc.add_heading("Switch CLI Capabilities", level=1)
doc.add_paragraph(
    "The switch CLI supports a practical access-switch configuration workflow. The tested switch was configured with hostname AN-1 and four VLANs, then assigned access and trunk behavior on individual interfaces."
)
add_table(doc,
    ["Command or command family", "Customer visible capability", "Test result"],
    [
        ("enable and configure terminal", "Move from user access to privileged and global configuration modes", "Accepted"),
        ("hostname and basic settings", "Set device identity and disable DNS lookup; configure password-encryption and enable-secret settings", "Accepted"),
        ("vlan <id> and name <name>", "Create and name VLANs", "Accepted for VLANs 80, 310, 20, and 190"),
        ("interface Fa0/1", "Select the uplink interface", "Accepted"),
        ("switchport trunk encapsulation dot1q", "Enable the modeled 802.1Q encapsulation flag", "Accepted"),
        ("switchport mode trunk", "Set the interface to trunk mode", "Accepted"),
        ("switchport trunk native vlan 190", "Set the native VLAN", "Accepted"),
        ("switchport trunk allowed vlan ...", "Restrict the trunk to a specified VLAN list", "Accepted for VLANs 20, 80, 190, and 310"),
        ("interface Fa0/2 through Fa0/4", "Assign access VLANs to edge ports", "Accepted"),
        ("show vlan brief", "Inspect the VLAN database and assignments", "Executed"),
        ("show interfaces trunk", "Inspect trunk mode, native VLAN, and allowed VLANs", "Executed"),
        ("copy running-config startup-config", "Save the current configuration", "Saved"),
    ], [1.85, 3.3, 1.6])

doc.add_heading("Verified Switch State", level=2)
add_bullet(doc, "Hostname and mode: AN-1 in Privileged EXEC")
add_bullet(doc, "VLAN database: 80 MANAGEMENT, 310 RECEPTION, 20 DATA, and 190 NATIVE")
add_bullet(doc, "Fa0/1: administratively up, 802.1Q trunk, native VLAN 190, allowed VLANs 20, 80, 190, and 310")
add_bullet(doc, "Fa0/2: administratively up, access VLAN 80")
add_bullet(doc, "Fa0/3: administratively up, access VLAN 310")
add_bullet(doc, "Fa0/4: administratively up, access VLAN 20")
add_bullet(doc, "Startup configuration: saved")

doc.add_heading("Router CLI Capabilities", level=1)
doc.add_paragraph(
    "The router CLI supports the core configuration path for a small routed training topology. The tested workflow reached Privileged EXEC, changed the hostname to CN-1, enabled the physical interface, applied interface addressing, configured an IP route, and saved the resulting configuration."
)
add_table(doc,
    ["Command or command family", "Customer visible capability", "Test result"],
    [
        ("enable and configure terminal", "Move from user access to privileged and global configuration modes", "Accepted"),
        ("hostname and basic settings", "Set device identity and disable DNS lookup; configure password-encryption and enable-secret settings", "Accepted"),
        ("interface Fa0/0", "Select and enable the physical router interface", "Accepted"),
        ("interface Fa0/0.<VLAN>", "Select router-on-a-stick subinterfaces", "Accepted during the test"),
        ("encapsulation dot1q <VLAN>", "Apply 802.1Q VLAN tagging to a subinterface", "Supported path exercised"),
        ("ip address <address> <mask>", "Apply an IP address and mask to an interface", "Supported path exercised"),
        ("ip route ...", "Configure a static route", "Accepted; one route present in final state"),
        ("show running-config", "Inspect the active configuration", "Executed"),
        ("show ip route", "Inspect route state", "Executed"),
        ("copy running-config startup-config", "Save the current configuration", "Saved"),
    ], [1.85, 3.3, 1.6])

doc.add_heading("Router Operating Limit", level=2)
doc.add_paragraph(
    "The router command path requires explicit exits when moving between interface contexts. In the source test run, subinterface retention was not treated as a clean pass. This means the router CLI is suitable for the current core training workflow, but subinterface-heavy procedures should be presented with explicit context exits and verified after configuration."
)

doc.add_heading("Customer Takeaway", level=1)
doc.add_paragraph(
    "The switch CLI currently provides a credible training model for access-switch configuration, VLAN administration, trunk setup, operational inspection, and configuration persistence. The router CLI provides the corresponding foundation for hostname and interface setup, routed addressing, 802.1Q subinterface work, static routes, inspection, and save operations."
)
doc.add_paragraph(
    "For training delivery, switch configuration procedures can be taught as a complete workflow. Router procedures should include an explicit context-management step and a final inspection step, especially when configuring subinterfaces."
)

doc.add_heading("Source and Scope", level=1)
doc.add_paragraph(
    "This report summarizes the verified command behavior and final states recorded in Router and Switch CLI Test Report dated 16 October 2026. It describes CLI capabilities and limitations only; development-environment details are outside the scope of this customer report."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Switch and Router CLI Capabilities | Customer Report")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(117, 117, 117)

doc.core_properties.title = "Customer Report Switch and Router CLI Capabilities"
doc.core_properties.subject = "Current tested switch and router CLI capabilities"
doc.core_properties.author = "Flight Management System"
doc.save(OUT)
print(OUT)
