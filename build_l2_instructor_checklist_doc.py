from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Unity\TSS2\Layer_2_Instructor_Required_Command_Checklist.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.size = Pt(size); styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
title_pPr = styles["Title"]._element.get_or_add_pPr()
title_border = title_pPr.find(qn("w:pBdr"))
if title_border is not None: title_pPr.remove(title_border)
styles["Title"].paragraph_format.space_after = Pt(4)
styles["Heading 1"].paragraph_format.space_before = Pt(13)
styles["Heading 1"].paragraph_format.space_after = Pt(5)
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(3)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table, color="D9D9D9"):
    tblPr = table._tbl.tblPr; border = tblPr.first_child_found_in("w:tblBorders")
    if border is None: border = OxmlElement("w:tblBorders"); tblPr.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = border.find(qn("w:" + edge))
        if el is None: el = OxmlElement("w:" + edge); border.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)

def cell_text(cell, text, bold=False, white=False, size=8.5):
    cell.text = ""; p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text)); r.font.name = "Aptos"; r.font.size = Pt(size); r.bold = bold
    if white: r.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i, h in enumerate(headers): cell_text(t.rows[0].cells[i], h, True, True); shade(t.rows[0].cells[i], "1F4E78")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value)
            if ri % 2: shade(cells[i], "F2F6FA")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullet(text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); p.add_run(text)

def code_block(text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F0F3F6"); pPr.append(shd)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.2)

title_p = doc.add_paragraph("Layer 2 Instructor Required Command Checklist", style="Title")
title_pPr = title_p._p.get_or_add_pPr()
for tag in ("w:pBdr", "w:bottom", "w:top", "w:left", "w:right"):
    el = title_pPr.find(qn(tag))
    if el is not None: title_pPr.remove(el)
p = doc.add_paragraph(); r = p.add_run("Student commands and instructor-checkable outcomes from Part 2 pages 9 through 15")
r.font.size = Pt(11); r.font.color.rgb = RGBColor(68, 68, 68); p.paragraph_format.space_after = Pt(9)

doc.add_paragraph("Purpose", style="Heading 1")
doc.add_paragraph("This is the narrowed Layer 2 scope for the TSS2 capstone. It lists what the student is required to perform on AN-1 and AN-2 and what the instructor is required to verify. VLAN IDs and most VLAN names are seeded in the student's Binder.")
doc.add_paragraph("The required flow is:")
bullet("Harden both switches.")
bullet("Create the five VLANs on each switch.")
bullet("Assign AN-1 access ports, trunks, and unused ports.")
bullet("Assign AN-2 access ports, trunk, and unused ports.")
bullet("Verify both switches and save both configurations.")

doc.add_paragraph("Device hardening on both switches", style="Heading 1")
code_block("enable\nconfigure terminal\nhostname AN-1\nenable secret <Binder enable secret>\nservice password-encryption\nno ip domain-lookup\nbanner motd #WARNING: DoD System#\n\nline console 0\n password <Binder console password>\n login\n exit\n\nline vty 0 4\n password <Binder VTY password>\n login\n exit")
doc.add_paragraph("Use hostname AN-2 when configuring the second switch.")
table(["Instructor check", "Required result"], [
    ("Identity", "Hostname is AN-1 or AN-2 as appropriate."),
    ("Privileged access", "Enable secret is configured; legacy enable password is not used."),
    ("Console and VTY", "Each line has the Binder password and login enabled."),
    ("Security services", "Password encryption is enabled and DNS lookup is disabled."),
    ("Banner", "The DoD warning banner is present with the required text."),
], [1.75, 5.15])

doc.add_paragraph("VLAN database on both switches", style="Heading 1")
code_block("vlan <MGMT ID>\n name <Management name>\n exit\nvlan <RECEP ID>\n name <Reception name>\n exit\nvlan <DATA ID>\n name <Data name>\n exit\nvlan <NATIVE ID>\n name <Native name>\n exit\nvlan <DEAD ID>\n name Dead_Vlan\n state suspend\n exit")
table(["Instructor check", "Required result"], [
    ("Database completeness", "All five VLANs exist on both AN-1 and AN-2."),
    ("Seeded values", "IDs and role names match the student's Binder exactly."),
    ("Unused VLAN", "The VLAN is named Dead_Vlan and is suspended."),
    ("Independence", "AN-2 has its own VLAN database; the trunk does not replicate VLAN configuration."),
], [1.75, 5.15])

doc.add_paragraph("AN-1 switchports", style="Heading 1")
doc.add_paragraph("Access ports", style="Heading 2")
code_block("interface fa0/2\n switchport mode access\n switchport access vlan <MGMT ID>\n exit\ninterface range fa0/3 - 4\n switchport mode access\n switchport access vlan <RECEP ID>\n exit\ninterface fa0/5\n switchport mode access\n switchport access vlan <MGMT ID>\n exit")
table(["Ports", "Required assignment"], [("Fa0/2", "Management"), ("Fa0/3 - Fa0/4", "Reception"), ("Fa0/5", "Management")], [1.75, 5.15])
doc.add_paragraph("Trunks", style="Heading 2")
code_block("interface fa0/1\n switchport mode trunk\n switchport trunk allowed vlan <MGMT>,<RECEP>,<DATA>\n switchport trunk native vlan <NATIVE ID>\n exit\ninterface fa0/6\n switchport mode trunk\n switchport trunk allowed vlan <MGMT>,<RECEP>,<DATA>\n switchport trunk native vlan <NATIVE ID>\n exit")
bullet("Fa0/1 carries VLANs between AN-1 and AN-2.")
bullet("Fa0/6 carries VLANs from AN-1 toward CN-1.")
bullet("The native VLAN is not included in the allowed VLAN list.")
doc.add_paragraph("Unused ports", style="Heading 2")
code_block("interface range fa0/7 - 8\n switchport mode access\n switchport access vlan <DEAD ID>\n shutdown\n exit")

doc.add_paragraph("AN-2 switchports", style="Heading 1")
doc.add_paragraph("Access ports", style="Heading 2")
code_block("interface fa0/2\n switchport mode access\n switchport access vlan <DATA ID>\n exit\ninterface fa0/3\n switchport mode access\n switchport access vlan <DATA ID>\n exit\ninterface fa0/4\n switchport mode access\n switchport access vlan <DATA ID>\n exit")
table(["Ports", "Required assignment"], [("Fa0/2", "Data"), ("Fa0/3", "Data"), ("Fa0/4", "Data")], [1.75, 5.15])
doc.add_paragraph("Trunk and unused ports", style="Heading 2")
code_block("interface fa0/1\n switchport mode trunk\n switchport trunk allowed vlan <MGMT>,<RECEP>,<DATA>\n switchport trunk native vlan <NATIVE ID>\n exit\ninterface range fa0/5 - 8\n switchport mode access\n switchport access vlan <DEAD ID>\n shutdown\n exit")
bullet("AN-2 Fa0/1 must match the AN-1 trunk allowed list and native VLAN.")
bullet("AN-2 Fa0/5 through Fa0/8 must be assigned to Dead_Vlan and shut down.")

doc.add_paragraph("Layer 2 verification and save", style="Heading 1")
code_block("show vlan brief\nshow interfaces trunk\ncopy running-config startup-config")
table(["Switch", "Required actions", "Instructor checks"], [
    ("AN-1", "Run both show commands and save.", "Five VLANs, correct port membership, Dead_Vlan ports down, both trunks correct, configuration saved."),
    ("AN-2", "Run both show commands and save.", "Five VLANs, correct port membership, Dead_Vlan ports down, matching trunk, configuration saved."),
], [1.0, 2.2, 3.8])
doc.add_paragraph("The instructor should require the student to read the output against the Binder, not merely execute the show commands.")

doc.add_paragraph("Not part of the Part 2 Layer 2 check", style="Heading 1")
doc.add_paragraph("These commands may exist in the broader simulator, but they are not required for the Layer 2 task described in the guide:")
code_block("ip address\nno ip address\nip route\nno ip route\nencapsulation dot1q\nswitchport trunk encapsulation dot1q\nshow ip route\nshow ip interface brief\nshow interfaces\nshow mac address-table\nping\nenable password")

doc.add_paragraph("Current implementation watch points", style="Heading 1")
bullet("The required command syntax is mostly present in the current processor.")
bullet("show vlan brief currently reports VLANs as active even when Dead_Vlan was suspended.")
bullet("show vlan brief does not clearly expose that Dead_Vlan ports are shut down.")
bullet("show interfaces trunk does not fully distinguish configured trunks from operationally active trunks.")
bullet("The current CLI accepts extra Layer 3 commands on switches, but those commands should not expand the Layer 2 evaluation scope.")

doc.add_paragraph("Source", style="Heading 2")
doc.add_paragraph("Net Capstone Instructor Guide (2), Part 2 - Layer 2, pages 9-15. Reviewed October 20 2026.")
doc.save(OUT)
print(OUT)
