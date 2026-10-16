from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\Unity\TSS2\Router_and_Switch_CLI_Test_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        new_row = table.add_row()
        set_row_cant_split(new_row)
        cells = new_row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12)):
    styles[name].font.name = "Aptos Display"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles[name].font.size = Pt(size)

title = doc.add_paragraph(style="Title")
title.add_run("Router and Switch CLI Test Report")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("Unity TSS2 training simulator | Play Mode validation | 16 October 2026")
run.italic = True
run.font.color.rgb = RGBColor(89, 89, 89)

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "The router and switch were installed in the Balanced_12_PIECE_LAB scenario and configured through the in-simulator CLI. "
    "The router configuration path reached Privileged EXEC with the expected hostname, route state, and saved startup configuration. "
    "After a clean Play Mode restart, the switch accepted both trunk commands; Fa0/1 became an 802.1Q trunk carrying the expected VLANs."
)

doc.add_heading("Test Environment", level=1)
add_table(doc, ["Item", "Value"], [
    ("Application", "Unity TSS2 training simulator"),
    ("Scene", "Balanced_12_PIECE_LAB"),
    ("Execution", "Unity Play Mode"),
    ("Project changes", "Added trunk-state regression assertions; no scene or prefab changes"),
    ("Console result", "0 errors and 0 warnings after testing"),
], widths=[1.8, 4.8])

doc.add_heading("Router Test", level=1)
doc.add_paragraph("Installed device: router at Rack01 U3. Runtime session key: Rack:Rack01:U03. Final hostname: CN-1.")

doc.add_heading("Router Commands and Effects", level=2)
router_rows = [
    ("enable", "User EXEC to Privileged EXEC", "Accepted"),
    ("configure terminal", "Privileged EXEC to Global Configuration", "Accepted"),
    ("hostname CN-1", "Changed device hostname to CN-1", "Accepted"),
    ("no ip domain-lookup", "Disabled DNS lookup", "Accepted"),
    ("service password-encryption", "Enabled password encryption setting", "Accepted"),
    ("enable secret training", "Configured the enable secret", "Accepted"),
    ("interface Fa0/0", "Selected the physical router interface", "Accepted"),
    ("no shutdown", "Enabled the selected interface", "Accepted"),
    ("interface Fa0/0.<VLAN>", "Selected router-on-a-stick subinterfaces", "Accepted during test"),
    ("encapsulation dot1q <VLAN>", "Applied 802.1Q VLAN tagging to subinterfaces", "Supported path exercised"),
    ("ip address <address> <mask>", "Applied interface addressing", "Supported path exercised"),
    ("ip route ...", "Configured a route; one route was present in final state", "Accepted"),
    ("show running-config / show ip route", "Read-only inspection and evidence capture", "Executed"),
    ("copy running-config startup-config", "Saved running configuration to startup state", "Saved"),
]
add_table(doc, ["Command", "Runtime state effect", "Result"], router_rows, widths=[2.25, 3.55, 0.8])

doc.add_heading("Router Final State", level=2)
add_bullet(doc, "Hostname: CN-1")
add_bullet(doc, "Mode: Privileged EXEC")
add_bullet(doc, "DNS lookup: disabled")
add_bullet(doc, "Routes: one route present")
add_bullet(doc, "Startup configuration: saved")
add_bullet(doc, "Qualification: the router command path was exercised, but the session required explicit interface exits for reliable movement between interface contexts; subinterface retention was not treated as a clean pass in that run.")

doc.add_heading("Switch Test", level=1)
doc.add_paragraph("Installed device: Cisco 2960 PoE Switch at Rack02 U3. Runtime session key: Rack:Rack02:U03. Final hostname: AN-1.")

doc.add_heading("Switch Commands and Effects", level=2)
switch_rows = [
    ("enable", "User EXEC to Privileged EXEC", "Accepted"),
    ("configure terminal", "Privileged EXEC to Global Configuration", "Accepted"),
    ("hostname AN-1", "Changed device hostname to AN-1", "Accepted"),
    ("no ip domain-lookup", "Disabled DNS lookup", "Accepted"),
    ("service password-encryption", "Enabled password encryption setting", "Accepted"),
    ("enable secret training", "Configured the enable secret", "Accepted"),
    ("vlan 80 / name MANAGEMENT", "Created and named VLAN 80", "Accepted"),
    ("vlan 310 / name RECEPTION", "Created and named VLAN 310", "Accepted"),
    ("vlan 20 / name DATA", "Created and named VLAN 20", "Accepted"),
    ("vlan 190 / name NATIVE", "Created and named VLAN 190", "Accepted"),
    ("interface Fa0/1", "Selected the uplink interface", "Accepted"),
    ("switchport trunk encapsulation dot1q", "Enabled the modeled 802.1Q encapsulation flag", "Accepted; dot1q=True"),
    ("switchport mode trunk", "Changed Fa0/1 to trunk mode", "Accepted; trunk=True"),
    ("switchport trunk native vlan 190", "Set native VLAN to 190", "Accepted"),
    ("switchport trunk allowed vlan 80,310,20,190", "Set four allowed VLANs", "Accepted"),
    ("no shutdown", "Enabled Fa0/1", "Accepted"),
    ("exit", "Returned from interface configuration to Global Configuration", "Required between interface selections"),
    ("interface Fa0/2 + access VLAN 80", "Enabled Fa0/2 and assigned VLAN 80", "Accepted"),
    ("interface Fa0/3 + access VLAN 310", "Enabled Fa0/3 and assigned VLAN 310", "Accepted"),
    ("interface Fa0/4 + access VLAN 20", "Enabled Fa0/4 and assigned VLAN 20", "Accepted"),
    ("end", "Returned to Privileged EXEC", "Accepted"),
    ("show vlan brief", "Read-only VLAN inspection and evidence capture", "Executed"),
    ("show interfaces trunk", "Read-only trunk inspection and evidence capture", "Executed"),
    ("copy running-config startup-config", "Saved running configuration to startup state", "Saved"),
]
add_table(doc, ["Command", "Runtime state effect", "Result"], switch_rows, widths=[2.55, 3.25, 0.8])

doc.add_heading("Switch Final State", level=2)
add_table(doc, ["Element", "Final state"], [
    ("Hostname and mode", "AN-1; Privileged EXEC"),
    ("VLAN database", "80 MANAGEMENT; 310 RECEPTION; 20 DATA; 190 NATIVE"),
    ("Fa0/1", "Administratively up; 802.1Q trunk; native VLAN 190; allowed VLANs 20, 80, 190, 310"),
    ("Fa0/2", "Administratively up; access VLAN 80"),
    ("Fa0/3", "Administratively up; access VLAN 310"),
    ("Fa0/4", "Administratively up; access VLAN 20"),
    ("Startup configuration", "Saved"),
], widths=[1.8, 4.8])

doc.add_heading("Conclusion", level=1)
doc.add_paragraph(
    "The access-switch configuration, VLAN database, interface VLAN assignments, inspection commands, and startup save behavior are working in Play Mode. "
    "The rerun also confirmed that switchport trunk encapsulation dot1q and switchport mode trunk are accepted after a clean Play Mode restart. "
    "Fa0/1 reported 802.1Q trunking with native VLAN 190 and all four expected VLANs associated."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("TSS2 CLI validation report")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(128, 128, 128)

doc.save(OUT)
print(OUT)
