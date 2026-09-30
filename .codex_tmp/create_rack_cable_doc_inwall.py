from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Unity\TSS2\Docs\Rack Cable Infrastructure Configuration Guide In Wall Update.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, top="120", left="120", bottom="120", right="120"):
    tc_pr = cell._tc.get_or_add_tcPr()
    margin = tc_pr.find(qn("w:tcMar"))
    if margin is None:
        margin = OxmlElement("w:tcMar")
        tc_pr.append(margin)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = margin.find(qn("w:" + side))
        if element is None:
            element = OxmlElement("w:" + side)
            margin.append(element)
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_padding(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9)
        if row_index == 0:
            set_repeat_table_header(row)
            for cell in row.cells:
                set_cell_shading(cell, "1F4E79")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        elif row_index % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F3F7FB")


def add_paragraph(doc, text="", style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    if text:
        run = paragraph.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
    return paragraph


def add_bullets(doc, items, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        add_paragraph(doc, item, style)


def add_numbered(doc, items):
    for item in items:
        add_paragraph(doc, item, "List Number")


def add_code_paragraph(doc, text):
    paragraph = add_paragraph(doc)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    return paragraph


def configure_styles(doc):
    styles = doc.styles
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)

    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Subtitle"].font.size = Pt(11)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 3"].font.size = Pt(11)
    styles["Heading 3"].font.bold = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.size = Pt(10.2)


def add_endpoint_table(doc):
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Endpoint source"
    hdr[1].text = "Endpoint ID pattern"
    hdr[2].text = "Meaning"

    rows = [
        ("Pre-mounted patch panel", "PatchPanel:Rack01:Port01", "A patch panel port that already exists in the open scene. Permanent infrastructure records should use these IDs when the patch panel is pre-installed."),
        ("Wallport", "WallPort:IT1:Port01", "A wall jack or room drop that exists in the scene. The permanent record links a patch panel port to one wallport endpoint as inferred in-wall infrastructure."),
        ("Player-installed rack equipment", "Rack:Rack01:U02:Port01", "An endpoint created when equipment is installed at runtime. Do not use this for pre-mounted patch panels."),
    ]
    for source, pattern, meaning in rows:
        cells = table.add_row().cells
        cells[0].text = source
        cells[1].text = pattern
        cells[2].text = meaning
    style_table(table)


def add_route_table(doc):
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Node type"
    hdr[1].text = "Typical placement"
    hdr[2].text = "Connect neighbors to"
    hdr[3].text = "Purpose"

    rows = [
        ("Rack rear node", "Behind each rack, near the patch panel and switch rear cable area", "Wall run node or nearby rack rear node", "Pull cable visuals into the rack path instead of straight through equipment."),
        ("Wall run node", "Along the wall, tray, or conduit path between rack and rooms", "Previous and next wall run nodes", "Creates clean horizontal or vertical cable runs."),
        ("Room drop node", "Near the point where a wall run enters a desk area or office", "Wall run node and wallport cluster node", "Separates building backbone routing from local desk routing."),
        ("Wallport cluster node", "Near a set of wall jacks or a specific wallport prefab", "Room drop node and adjacent wallport cluster node if needed", "Keeps visible player-created cable segments short and visually aligned with wallports."),
        ("Desk node", "Behind a desk or equipment placement area", "Wallport cluster node", "Guides player-created cables between wallport endpoints and placed desktop/printer equipment."),
    ]
    for node_type, placement, neighbors, purpose in rows:
        cells = table.add_row().cells
        cells[0].text = node_type
        cells[1].text = placement
        cells[2].text = neighbors
        cells[3].text = purpose
    style_table(table)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Rack Cable Infrastructure Configuration Guide")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Permanent patch panel records endpoint checks and route node setup for the TSS2 Unity scene")

    add_paragraph(
        doc,
        "This guide explains how to finish the rack cable infrastructure work in the TSS2 training scene. "
        "It focuses on permanent patch-panel-to-wallport records, endpoint ID verification, and manual route node placement for visible cable runs. "
        "The goal is a deterministic training setup where pre-mounted patch panels and wallports resolve cleanly, while hidden building infrastructure remains represented as state rather than scene geometry.",
    )

    doc.add_heading("Permanent Infrastructure Records", level=1)
    add_paragraph(
        doc,
        "A permanent infrastructure record represents a building cable that already exists before the student begins the exercise. "
        "In this project, the record links one pre-mounted patch panel port to one wallport endpoint. It is not an inventory cable, should not reduce the player cable count, should not be removable through normal cable interaction, and should not spawn a visible cable. The player should infer that this cable is installed inside the walls, floors, or ceilings of the building."
    )
    add_paragraph(
        doc,
        "For the current scene, the record belongs in the scenario asset rather than in rack equipment runtime logic. "
        "Use this only for infrastructure that should exist immediately when the scenario starts."
    )

    doc.add_heading("How to Add Remaining Records", level=2)
    add_numbered(
        doc,
        [
            "Open Assets/Data/TSS/Scenarios/Balanced_12_Piece_Lab.asset in the Unity Inspector.",
            "Find the Permanent Connections array on the TssScenarioDefinition component.",
            "Increase the array size for each additional building cable you want to preconfigure.",
            "For each row, enter a stable connection ID such as Infrastructure002 or PatchRack01Port02ToOffice01Port02.",
            "Assign the cable asset, usually Cable_StraightThrough, unless the lab intentionally needs another cable type.",
            "Set cableType to match the selected cable asset.",
            "Set patchPanelEndpointId to the pre-mounted patch panel endpoint ID, for example PatchPanel:Rack01:Port02.",
            "Set wallportEndpointId to the matching wallport endpoint ID, for example WallPort:Office01:Port02.",
            "Enter Play Mode or start the scenario and confirm the connection appears in the Cable Connections HUD under Permanent Infrastructure.",
        ],
    )

    doc.add_heading("How to Find or Establish the Endpoint IDs", level=3)
    add_paragraph(
        doc,
        "Before adding a row to the Permanent Connections array, identify the exact endpoint ID string for both sides of the inferred building cable. The scenario record does not point to a GameObject directly. It stores text IDs, so a typo or duplicate ID will make the record fail to resolve or resolve to the wrong port."
    )
    add_paragraph(
        doc,
        "Start with the patch panel side. In the Hierarchy, select the pre-mounted rack patch panel, expand its port objects, and inspect each TssPortEndpoint component. The field to copy is endpointId. For pre-mounted patch panels, use the scene-authored PatchPanel prefix, such as PatchPanel:Rack01:Port01. Do not use a Rack:Rack01:U02 style ID unless the patch panel is installed at runtime by rack install logic."
    )
    add_paragraph(
        doc,
        "Then check the wallport side. Select the wallport object in the room or desk area, expand its individual port objects, and inspect each TssPortEndpoint component. Copy the wallport endpointId, such as WallPort:IT1:Port01 or WallPort:Office01:Port02. This is the ID that represents the visible wall jack where the player expects the hidden building cable to terminate."
    )
    add_paragraph(
        doc,
        "If the endpoint ID is missing, vague, or duplicated, establish it on the endpoint object before adding the permanent record. Use a stable pattern that names the physical system, location, and port number. After changing an endpoint ID, update any scenario records that reference the old value."
    )
    add_bullets(
        doc,
        [
            "Patch panel endpoint field: copy into patchPanelEndpointId.",
            "Wallport endpoint field: copy into wallportEndpointId.",
            "Connection ID field: create a human-readable record name such as Infrastructure002 or Rack01Port02ToIT1Port02.",
            "Verification: enter Play Mode and confirm the HUD row appears under Permanent Infrastructure without creating a visible cable.",
        ],
    )

    doc.add_heading("Endpoint ID Examples for Permanent Records", level=3)
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Permanent record field"
    hdr[1].text = "Example value"
    hdr[2].text = "Where to find or set it"
    hdr[3].text = "Common mistake"
    rows = [
        (
            "connectionId",
            "Infrastructure002",
            "Create this in the scenario row. It only needs to be stable and readable.",
            "Reusing the same connectionId for multiple records.",
        ),
        (
            "patchPanelEndpointId",
            "PatchPanel:Rack01:Port02",
            "Copy from the TssPortEndpoint on the pre-mounted patch panel port.",
            "Using Rack:Rack01:U02:Port02 for a panel that already exists in the scene.",
        ),
        (
            "wallportEndpointId",
            "WallPort:IT1:Port02",
            "Copy from the TssPortEndpoint on the wallport port object.",
            "Guessing a room name instead of copying the exact endpointId.",
        ),
        (
            "cable and cableType",
            "Cable_StraightThrough and StraightThrough",
            "Use the scenario cable asset and matching cableType.",
            "Choosing a type unsupported by either endpoint.",
        ),
    ]
    for field, example, source, mistake in rows:
        cells = table.add_row().cells
        cells[0].text = field
        cells[1].text = example
        cells[2].text = source
        cells[3].text = mistake
    style_table(table)

    add_paragraph(
        doc,
        "Do not use runtime install endpoint IDs for pre-mounted patch panels. IDs like Rack:Rack01:U02:Port01 are created for equipment installed by the player or by rack install logic. Pre-mounted patch panels in the scene should use their scene-authored IDs, such as PatchPanel:Rack01:Port01."
    )
    add_endpoint_table(doc)

    doc.add_heading("Permanent Record Validation Checklist", level=2)
    add_bullets(
        doc,
        [
            "The patch panel endpoint ID exists in the open scene before Play Mode starts.",
            "The wallport endpoint ID exists in the open scene before Play Mode starts.",
            "The two endpoint IDs are not already used by another permanent or player-created connection.",
            "The cable type is supported by both endpoint definitions.",
            "The HUD shows the row under Permanent Infrastructure after scenario start.",
            "No visible cable object is created for the patch-panel-to-wallport record.",
            "Selecting either endpoint without a held cable does not remove the permanent connection.",
        ],
    )

    doc.add_heading("Endpoint ID Uniqueness", level=1)
    add_paragraph(
        doc,
        "Every selectable port needs a unique endpoint ID because cable records store endpoint IDs, not object references. "
        "If two patch panel ports share the same ID, a permanent connection can resolve to the wrong visual port, block the wrong port, or appear duplicated in the HUD."
    )

    doc.add_heading("What to Verify", level=2)
    add_bullets(
        doc,
        [
            "Each pre-mounted patch panel port has one unique endpoint ID.",
            "The ID describes the physical endpoint clearly enough to troubleshoot later.",
            "The ID pattern is consistent across panels and racks.",
            "No two active scene endpoints share the same final ID.",
        ],
    )

    doc.add_heading("Recommended Naming Pattern", level=2)
    add_paragraph(
        doc,
        "For pre-mounted patch panels, use a scene-stable pattern that names the system, rack, and port number."
    )
    add_code_paragraph(doc, "PatchPanel:Rack01:Port01")
    add_code_paragraph(doc, "PatchPanel:Rack01:Port02")
    add_code_paragraph(doc, "PatchPanel:Rack02:Port01")
    add_paragraph(
        doc,
        "If there are two physical patch panels in the same rack, include a panel identifier before the port number, such as PatchPanel:Rack01:PanelA:Port01 and PatchPanel:Rack01:PanelB:Port01. The exact wording matters less than uniqueness and consistency."
    )

    doc.add_heading("How to Check Uniqueness in Unity", level=2)
    add_numbered(
        doc,
        [
            "In the Hierarchy, locate the pre-mounted patch panel objects.",
            "Expand the panel prefab or scene object until the individual port objects are visible.",
            "Select each port object and inspect its TssPortEndpoint component.",
            "Confirm endpointId is populated and follows the chosen naming pattern.",
            "Use a temporary editor check or manual search to list all TssPortEndpoint components and sort by endpointId.",
            "Correct any duplicate IDs before adding permanent records that reference those ports.",
        ],
    )

    add_paragraph(
        doc,
        "When fixing duplicates, change the endpoint ID on the scene or prefab source that owns the port. Avoid only changing the scenario record to work around a duplicate, because the duplicate will still affect cable selection and connection blocking later."
    )

    doc.add_heading("Cable Route Node Setup", level=1)
    add_paragraph(
        doc,
        "Route nodes are manual waypoints that guide visible cable visuals through believable paths. They are for cables the player can see, such as rack-side patching or player-created endpoint cables. They are not needed for patch-panel-to-wallport permanent infrastructure records, because those records represent in-wall building wiring and should not create visible TssCableVisual geometry."
    )

    add_route_table(doc)

    doc.add_heading("Step 1 Place Route Nodes Behind Each Rack", level=2)
    add_paragraph(
        doc,
        "Create one or more TssCableRouteNode objects behind each rack where cables should visually enter or leave the rack. Place the node near the rear patch field or cable-management area, not inside the rack model. The node should be close enough that a cable starting from a patch panel port looks like it bends into the rack cable path."
    )
    add_bullets(
        doc,
        [
            "Use one rack rear node for a simple rack.",
            "Use separate upper and lower rear nodes if the rack has visibly different cable paths.",
            "Name nodes clearly, for example Route_Rack01_RearUpper or Route_Rack01_RearLower.",
        ],
    )

    doc.add_heading("Step 2 Place Wall Run Nodes Along Cable Paths", level=2)
    add_paragraph(
        doc,
        "Wall run nodes describe the main path from the rack area to each room or wallport cluster. Place them where a cable tray, conduit, ceiling path, or wall route would realistically run. Keep the path simple and readable; the goal is clean training visualization, not full building wiring simulation."
    )
    add_bullets(
        doc,
        [
            "Use enough nodes to turn corners cleanly.",
            "Avoid placing nodes in the middle of rooms unless the cable should visibly cross that area.",
            "Keep node height consistent along a run so cable segments do not zigzag vertically.",
        ],
    )

    doc.add_heading("Step 3 Place Nodes Near Desk and Wallport Clusters", level=2)
    add_paragraph(
        doc,
        "At the destination end, place a node near each group of wallports or desk drops only for visible player-created cables. This gives a cable a clean final approach to the endpoint when the student connects a wallport to a desktop, printer, switch, or other visible device. For rooms with multiple wallports, one cluster node can serve several ports if the ports are close together."
    )
    add_bullets(
        doc,
        [
            "Use one node per wallport cluster for most rooms.",
            "Use additional desk nodes when player-placed desktop or printer cables should route behind furniture.",
            "Place the final node close enough to the endpoint that the last segment is short and visually intentional.",
        ],
    )

    doc.add_heading("Step 4 Connect Neighbors Only to Adjacent Nodes", level=2)
    add_paragraph(
        doc,
        "Each TssCableRouteNode has a neighbors array. Add only the nodes that are directly reachable from that node. Do not connect every node to every other node. Dense neighbor lists make the route graph harder to reason about and can produce odd shortcuts."
    )
    add_paragraph(doc, "A clean rack-to-room path should look like this:")
    add_code_paragraph(doc, "Rack rear node -> wall run node -> room drop node -> wallport cluster node")
    add_paragraph(doc, "For each pair, set the neighbor relationship in both directions unless the route should intentionally be one way.")

    doc.add_heading("Step 5 Keep Route Nodes Slightly Clear of Geometry", level=2)
    add_paragraph(
        doc,
        "Place nodes slightly above the floor, behind furniture, or along the intended tray height so cable cylinders do not clip through desks, walls, or rack doors. Use the project unit convention: one Unity unit is one meter. Small offsets matter; moving a node by 0.05 to 0.15 meters can prevent visible clipping."
    )
    add_bullets(
        doc,
        [
            "For floor-level paths, start around y = 0.08 to 0.15 meters.",
            "For wall or tray paths, use the visible cable tray or wall route height.",
            "Keep endpoint-to-node segments short around crowded equipment.",
        ],
    )

    doc.add_heading("Step 6 Leave Direct Fallback Where It Is Acceptable", level=2)
    add_paragraph(
        doc,
        "Not every possible visible cable needs a complete route on day one. If no route provider or node path is available, TssCableVisual still draws a direct segment between endpoints. That fallback is useful for unfinished areas, temporary testing, and low-priority visible cable runs."
    )
    add_paragraph(
        doc,
        "Use route nodes first for high-visibility runs: rack-side patching, wallport to desk, wallport to printer, and any player-visible cable path that otherwise crosses through major scenery. Do not add visible route nodes solely to render hidden patch-panel-to-wallport infrastructure."
    )

    doc.add_heading("Route Node Validation Checklist", level=2)
    add_bullets(
        doc,
        [
            "Permanent infrastructure records appear in state and HUD without visible scene cable geometry.",
            "Player-created cables still draw correctly when a route exists.",
            "Fallback direct segments still work where no route exists.",
            "No cable segment clips heavily through rack faces, desks, walls, or floors.",
            "The graph remains sparse and readable in the Inspector.",
            "Node names make it obvious which rack, wall run, or room they belong to.",
        ],
    )

    doc.add_heading("Practical Setup Order", level=1)
    add_numbered(
        doc,
        [
            "Confirm permanent records resolve immediately for pre-mounted patch panels and wallports.",
            "Fix duplicate endpoint IDs before adding more scenario records.",
            "Add route nodes for visible rack-side or player-created cable paths only.",
            "Verify the Cable Connections HUD shows the Permanent Infrastructure section.",
            "Connect and disconnect a player cable to confirm player cable behavior still works separately.",
            "Expand route nodes only for cable paths that look wrong in Play Mode.",
        ],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
