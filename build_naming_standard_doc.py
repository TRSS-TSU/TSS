from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Unity\TSS2\Naming_Standard_Tab_Review.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def make_table(doc, floor, rack, rows):
    doc.add_heading(f'{floor} Floor Equipment', level=2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    widths = [Inches(0.8), Inches(1.55), Inches(3.8)]
    headers = ['Order', 'Equipment Name', 'Mounting / Naming Notes']
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_text(cell, h, bold=True, color=(255, 255, 255))
        shade(cell, '1F4E78')
    for n, (name, note) in enumerate(rows, 1):
        cells = table.add_row().cells
        for i, width in enumerate(widths):
            cells[i].width = width
        set_cell_text(cells[0], str(n), bold=True)
        set_cell_text(cells[1], name, bold=True)
        set_cell_text(cells[2], note)
        if n % 2 == 0:
            for cell in cells:
                shade(cell, 'EAF2F8')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f'Rack: {rack}')
    r.italic = True
    r.font.color.rgb = RGBColor(89, 89, 89)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for style_name, size in [('Title', 22), ('Heading 1', 16), ('Heading 2', 12)]:
    style = doc.styles[style_name]
    style.font.name = 'Aptos'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph('Naming Standard', style='Title')
doc.add_paragraph('Rack Equipment Mounting Requirements', style='Subtitle')
doc.add_heading('Equipment Mounting Requirements', level=1)
p = doc.add_paragraph()
p.add_run('Mount equipment in each network rack from top to bottom in the order shown below. ').bold = True
p.add_run('Use the exact equipment names in the tables when labeling rack components, documenting connections, or referring to equipment in the training scenario. Keep the rack order consistent so the physical installation can be checked quickly against the naming standard.')

make_table(doc, 'First', 'RACK01', [
    ('PDU1', 'Power distribution unit for the ground-floor rack.'),
    ('Patch1', 'Patch panel for the ground-floor rack.'),
    ('AN-1', 'Access switch for the ground-floor rack.'),
    ('CN-1', 'Router for the ground-floor rack.'),
    ('IT2', 'Server; retain the exact name IT2.'),
])

make_table(doc, 'Second', 'RACK02', [
    ('PDU2', 'Power distribution unit for the second-floor rack.'),
    ('Patch2', 'Patch panel for the second-floor rack.'),
    ('AN-2', 'Access switch for the second-floor rack.'),
])

doc.add_heading('Server and Previous Installation', level=1)
doc.add_paragraph('IT2 is the server mounted in the ground-floor RACK01 installation. The server was installed previously; preserve its existing placement and established connections unless the scenario explicitly directs a change. When documenting or referring to it, use the exact name IT2.')

doc.add_paragraph('Source: live Naming Standard / Rack Build Order tab in the TSS2 Unity project.', style='Subtitle')
doc.save(OUT)
print(OUT)
